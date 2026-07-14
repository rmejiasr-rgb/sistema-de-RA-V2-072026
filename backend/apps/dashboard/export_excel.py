"""
Exportación a Excel (§8): hoja de detalle + hoja de resumen con conteos.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.uploads.models import ResultadoCriterio

AZUL = "FF003876"
NARANJA = "FFF37021"
VERDE = "FF1F8A4C"
AMARILLO = "FFC98A12"
ROJO = "FFC0392B"

HEADER_FILL = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)


def _estilar_encabezado(ws, fila, n_columnas):
    for col in range(1, n_columnas + 1):
        celda = ws.cell(row=fila, column=col)
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center")


def _autoancho(ws, max_ancho=50):
    for col in ws.columns:
        largo = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(largo + 2, max_ancho)


def exportar_evaluaciones_excel(evaluaciones, umbral_verde, umbral_amarillo):
    """Genera un .xlsx con hoja Resumen (una fila por evaluación) y hoja Detalle
    (una fila por estudiante×evaluación). Devuelve bytes."""
    wb = Workbook()

    # --- Hoja Resumen ---
    ws = wb.active
    ws.title = "Resumen"
    encabezados = [
        "Periodo", "Escuela", "Materia", "Sección", "Profesor", "RA", "Nivel",
        "Actividad", "Fecha", "% Declarado", "% Recalculado", "Semáforo", "Versión",
    ]
    ws.append(encabezados)
    _estilar_encabezado(ws, 1, len(encabezados))

    evaluaciones = list(
        evaluaciones.select_related("materia", "materia__escuela", "ra", "periodo", "carga")
    )
    for ev in evaluaciones:
        pct = float(ev.pct_aprobados_recalculado)
        if pct >= umbral_verde:
            semaforo, color = "Verde", VERDE
        elif pct >= umbral_amarillo:
            semaforo, color = "Amarillo", AMARILLO
        else:
            semaforo, color = "Rojo", ROJO
        ws.append([
            ev.periodo.codigo, ev.materia.escuela.codigo, ev.materia.codigo, ev.seccion,
            ev.profesor_nombre, ev.ra.codigo, ev.nivel, ev.actividad,
            ev.fecha_actividad.isoformat(), float(ev.pct_aprobados_declarado),
            pct, semaforo, ev.carga.version,
        ])
        celda_sem = ws.cell(row=ws.max_row, column=12)
        celda_sem.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        celda_sem.font = Font(color="FFFFFFFF", bold=True)
    _autoancho(ws)
    ws.freeze_panes = "A2"

    # --- Hoja Detalle ---
    wd = wb.create_sheet("Detalle")
    encabezados_d = ["Periodo", "Materia", "Sección", "RA", "Grupo", "Estudiante",
                     "Carrera", "Nota (%)"]
    wd.append(encabezados_d)
    _estilar_encabezado(wd, 1, len(encabezados_d))

    eval_ids = [ev.id for ev in evaluaciones]
    from apps.uploads.models import ResultadoEstudiante
    estudiantes = (
        ResultadoEstudiante.objects.filter(evaluacion_id__in=eval_ids)
        .select_related("evaluacion__materia", "evaluacion__ra", "evaluacion__periodo")
        .order_by("evaluacion__materia__codigo", "evaluacion__ra__codigo", "nombre")
    )
    for est in estudiantes:
        ev = est.evaluacion
        wd.append([
            ev.periodo.codigo, ev.materia.codigo, ev.seccion, ev.ra.codigo,
            est.grupo, est.nombre, est.carrera, float(est.nota),
        ])
    _autoancho(wd)
    wd.freeze_panes = "A2"

    # --- Hoja Conteos por criterio ---
    wc = wb.create_sheet("Conteos criterio")
    wc.append(["Criterio", "Excelente", "Satisfactorio", "Insatisfactorio", "Total", "% Aprobados"])
    _estilar_encabezado(wc, 1, 6)
    from django.db.models import Count
    agregados = {}
    for fila in (ResultadoCriterio.objects
                 .filter(resultado_estudiante__evaluacion_id__in=eval_ids)
                 .values("criterio__nombre", "nivel")
                 .annotate(total=Count("id"))):
        nombre = fila["criterio__nombre"]
        agregados.setdefault(nombre, {"EXCELENTE": 0, "SATISFACTORIO": 0, "INSATISFACTORIO": 0})
        agregados[nombre][fila["nivel"]] = fila["total"]
    for nombre in sorted(agregados):
        d = agregados[nombre]
        total = d["EXCELENTE"] + d["SATISFACTORIO"] + d["INSATISFACTORIO"]
        pct = round((d["EXCELENTE"] + d["SATISFACTORIO"]) / total * 100, 2) if total else 0
        wc.append([nombre, d["EXCELENTE"], d["SATISFACTORIO"], d["INSATISFACTORIO"], total, pct])
    _autoancho(wc)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
