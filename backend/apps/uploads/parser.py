"""
Parser de archivos Excel de rúbricas de Resultados de Aprendizaje (RA) — UNIMET.

Ver especificación §5. Localiza los datos por celdas ancla y por búsqueda
dinámica de encabezados (el número de criterios y la posición de algunas
columnas varían de un archivo a otro).
"""

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl

TRES_DECIMALES = Decimal("0.001")


def _a_decimal(valor, defecto=None):
    if valor is None:
        return defecto
    try:
        return Decimal(str(valor)).quantize(TRES_DECIMALES, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return defecto

NIVELES_VALIDOS = {"EXCELENTE", "SATISFACTORIO", "INSATISFACTORIO"}

NOMBRE_ARCHIVO_RE = re.compile(
    r"^(?P<codmat>[A-Z0-9]+)-(?P<seccion>[A-Z0-9]+)-(?P<profesor>[A-ZÀ-ÖØ-öø-ÿ_]+)-"
    r"(?P<periodo>\d{4}-\d+)-(?P<ra>RA\d+(?:-[A-Z]+)*)-N(?P<nivel>\d+)\.xlsx$",
    re.IGNORECASE,
)


class ParserError(Exception):
    """Error de estructura (V1): el archivo no parece una plantilla RA UNIMET válida."""

    def __init__(self, mensaje, codigo="V1_ESTRUCTURA"):
        self.mensaje = mensaje
        self.codigo = codigo
        super().__init__(mensaje)


@dataclass
class EstudianteParseado:
    fila_excel: int
    grupo: str
    nombre: str
    carrera: str
    nota: Decimal
    niveles_por_criterio: dict = field(default_factory=dict)


@dataclass
class ArchivoParseado:
    asignatura_texto: str
    seccion: str
    profesor: str
    periodo_codigo: str
    actividad: str
    fecha_actividad: date
    ra_codigo: str
    nivel: int
    pct_aprobados_declarado: Decimal
    criterios: list
    estudiantes: list
    nombre_archivo_info: dict | None


def normalizar_texto(valor):
    """Mayúsculas, sin acentos, espacios colapsados. Para comparaciones tolerantes."""
    if valor is None:
        return ""
    texto = str(valor).strip()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto)
    return texto.upper().strip()


def normalizar_trimestre(valor):
    """Normaliza el Trimestre (E3) a texto 'AAAA-T'.

    Excel suele convertir '2526-3' a datetime(2526, 3, 1) porque lo interpreta
    como fecha. Si ya viene como texto con el patrón correcto, se respeta.
    """
    if isinstance(valor, datetime):
        return f"{valor.year}-{valor.month}"
    texto = str(valor).strip()
    if re.fullmatch(r"\d{4}-\d{1,2}", texto):
        return texto
    raise ParserError(
        f"El Trimestre (celda E3) tiene un valor no reconocible: '{valor}'. "
        "Se esperaba un código de periodo tipo 'AAAA-T' (ej. 2526-3).",
        codigo="V1_TRIMESTRE",
    )


def normalizar_ra_codigo(valor):
    """'RA10 IP UNIMET' -> 'RA10-IP'; 'RA8 UNIMET' -> 'RA8'."""
    if not valor:
        raise ParserError("La celda de RA (C7) está vacía.", codigo="V1_RA")
    texto = normalizar_texto(valor)
    texto = re.sub(r"\bUNIMET\b", "", texto).strip()
    texto = re.sub(r"\s+", "-", texto)
    if not re.match(r"^RA\d+", texto):
        raise ParserError(
            f"El valor del RA (celda C7) no tiene el formato esperado: '{valor}'.",
            codigo="V1_RA",
        )
    return texto


def extraer_nivel(valor):
    """'Nivel 2' -> 2."""
    if valor is None:
        raise ParserError("La celda de Nivel (D7) está vacía.", codigo="V1_NIVEL")
    match = re.search(r"(\d+)", str(valor))
    if not match:
        raise ParserError(
            f"El valor del Nivel (celda D7) no contiene un número: '{valor}'.",
            codigo="V1_NIVEL",
        )
    return int(match.group(1))


def normalizar_nivel_criterio(valor):
    texto = normalizar_texto(valor)
    if texto in NIVELES_VALIDOS:
        return texto
    return None


def parsear_nombre_archivo(nombre_archivo):
    """Extrae (codmat, seccion, profesor, periodo, ra, nivel) del nombre del archivo.

    Devuelve None si el nombre no sigue el patrón esperado (§5): la validación
    cruzada V2 lo maneja como discrepancia, no aquí.
    """
    match = NOMBRE_ARCHIVO_RE.match(nombre_archivo.strip())
    if not match:
        return None
    datos = match.groupdict()
    return {
        "codmat": datos["codmat"].upper(),
        "seccion": datos["seccion"],
        "profesor": datos["profesor"].replace("_", " ").strip(),
        "periodo": datos["periodo"],
        "ra": datos["ra"].upper(),
        "nivel": int(datos["nivel"]),
    }


def _celda(ws, coordenada):
    return ws[coordenada].value


def _localizar_columna_resultado(ws, fila_busqueda=2, max_col=20):
    for col in range(1, max_col + 1):
        valor = ws.cell(row=fila_busqueda, column=col).value
        if valor and normalizar_texto(valor) == "RESULTADO":
            return col
    return None


def _localizar_fila_encabezado_estudiantes(ws, max_filas=120, max_col=25):
    """Busca la fila de encabezado de la tabla de detalle de estudiantes.

    Estrategia primaria (spec): columna B == 'Nombre y Apellido'.
    Estrategia de respaldo (archivos reales observados con etiqueta mal escrita):
    la fila donde alguna celda contiene 'Nota', que siempre encabeza la última
    columna de la tabla de detalle.
    """
    candidata_por_nota = None
    for row in range(1, max_filas + 1):
        valor_b = ws.cell(row=row, column=2).value
        if valor_b and normalizar_texto(valor_b) == "NOMBRE Y APELLIDO":
            return row
        for col in range(3, max_col + 1):
            valor = ws.cell(row=row, column=col).value
            if valor and "NOTA" in normalizar_texto(valor) and valor_b:
                candidata_por_nota = row
                break
    return candidata_por_nota


def parse_workbook(archivo, nombre_archivo):
    """Parsea un archivo Excel de rúbrica RA. Lanza ParserError (V1) si la
    estructura no corresponde a una plantilla RA UNIMET."""
    try:
        wb = openpyxl.load_workbook(archivo, data_only=True)
    except Exception as exc:
        raise ParserError(
            f"No fue posible abrir el archivo como Excel válido ({exc}).",
            codigo="V1_ARCHIVO",
        )

    ws = wb[wb.sheetnames[0]]

    asignatura = _celda(ws, "C2")
    seccion = _celda(ws, "E2")
    profesor = _celda(ws, "C3")
    trimestre_raw = _celda(ws, "E3")
    actividad = _celda(ws, "C4")
    fecha_actividad_raw = _celda(ws, "C5")
    ra_raw = _celda(ws, "C7")
    nivel_raw = _celda(ws, "D7")

    faltantes = []
    if not asignatura:
        faltantes.append("Asignatura (C2)")
    if not seccion:
        faltantes.append("Sección (E2)")
    if not profesor:
        faltantes.append("Profesor (C3)")
    if not trimestre_raw:
        faltantes.append("Trimestre (E3)")
    if not actividad:
        faltantes.append("Actividad (C4)")
    if not fecha_actividad_raw:
        faltantes.append("Fecha de la Actividad (C5)")
    if not ra_raw:
        faltantes.append("Objetivo educativo / RA (C7)")
    if not nivel_raw:
        faltantes.append("Nivel (D7)")
    if faltantes:
        raise ParserError(
            "El archivo no parece una plantilla RA UNIMET: faltan celdas ancla "
            f"esperadas: {', '.join(faltantes)}.",
            codigo="V1_ESTRUCTURA",
        )

    seccion = str(seccion).strip()
    profesor = str(profesor).strip()
    actividad = str(actividad).strip()
    periodo_codigo = normalizar_trimestre(trimestre_raw)
    ra_codigo = normalizar_ra_codigo(ra_raw)
    nivel = extraer_nivel(nivel_raw)

    if isinstance(fecha_actividad_raw, datetime):
        fecha_actividad = fecha_actividad_raw.date()
    else:
        raise ParserError(
            f"La Fecha de la Actividad (C5) no es una fecha válida: '{fecha_actividad_raw}'.",
            codigo="V1_FECHA",
        )

    col_resultado = _localizar_columna_resultado(ws)
    if col_resultado is None:
        raise ParserError(
            "No se encontró la etiqueta 'Resultado' en la fila 2: no es posible "
            "ubicar el % de aprobados declarado.",
            codigo="V1_RESULTADO",
        )
    pct_raw = ws.cell(row=4, column=col_resultado).value
    if pct_raw is None:
        raise ParserError(
            "La celda del % de aprobados (fila 4, bajo la columna 'Resultado') está vacía.",
            codigo="V1_RESULTADO",
        )
    try:
        pct_fraccion = Decimal(str(pct_raw))
    except InvalidOperation:
        raise ParserError(
            f"El % de aprobados declarado no es numérico: '{pct_raw}'.", codigo="V1_RESULTADO"
        )
    pct_aprobados_declarado = (pct_fraccion * 100).quantize(TRES_DECIMALES, rounding=ROUND_HALF_UP)

    fila_encabezado = _localizar_fila_encabezado_estudiantes(ws)
    if fila_encabezado is None:
        raise ParserError(
            "No se encontró la tabla de 'Resultados detallados' por estudiante "
            "(se buscó la columna 'Nombre y Apellido' / encabezado con 'Nota').",
            codigo="V1_DETALLE",
        )

    # Columna donde está "Nota" en la fila de encabezado -> límite de criterios
    col_nota = None
    for col in range(3, 30):
        valor = ws.cell(row=fila_encabezado, column=col).value
        if valor and "NOTA" in normalizar_texto(valor):
            col_nota = col
            break
    if col_nota is None or col_nota < 4:
        raise ParserError(
            "No se encontró la columna 'Nota (%)' en el encabezado de la tabla de detalle.",
            codigo="V1_DETALLE",
        )

    criterios = []
    for col in range(4, col_nota):
        nombre_criterio = ws.cell(row=fila_encabezado, column=col).value
        if nombre_criterio and str(nombre_criterio).strip():
            criterios.append(str(nombre_criterio).strip())
    if not criterios:
        raise ParserError(
            "No se detectaron columnas de criterios entre la columna de Carrera y Nota (%).",
            codigo="V1_DETALLE",
        )

    estudiantes = []
    nombres_vistos = set()
    row = fila_encabezado + 1
    while True:
        nombre = ws.cell(row=row, column=2).value
        if not nombre or not str(nombre).strip():
            break
        nombre = str(nombre).strip()

        grupo_raw = ws.cell(row=row, column=1).value
        if isinstance(grupo_raw, float) and grupo_raw.is_integer():
            grupo = str(int(grupo_raw))
        elif grupo_raw is None:
            grupo = ""
        else:
            grupo = str(grupo_raw).strip()

        carrera = ws.cell(row=row, column=3).value
        carrera = str(carrera).strip() if carrera else ""

        nota_raw = ws.cell(row=row, column=col_nota).value
        nota = _a_decimal(nota_raw)

        niveles = {}
        for idx, nombre_criterio in enumerate(criterios):
            col = 4 + idx
            valor_nivel = ws.cell(row=row, column=col).value
            niveles[nombre_criterio] = valor_nivel

        estudiantes.append(
            EstudianteParseado(
                fila_excel=row,
                grupo=grupo,
                nombre=nombre,
                carrera=carrera,
                nota=nota,
                niveles_por_criterio=niveles,
            )
        )
        nombres_vistos.add(normalizar_texto(nombre))
        row += 1

    if not estudiantes:
        raise ParserError(
            "No se encontraron estudiantes en la tabla de 'Resultados detallados'.",
            codigo="V1_DETALLE",
        )

    return ArchivoParseado(
        asignatura_texto=str(asignatura).strip(),
        seccion=seccion,
        profesor=profesor,
        periodo_codigo=periodo_codigo,
        actividad=actividad,
        fecha_actividad=fecha_actividad,
        ra_codigo=ra_codigo,
        nivel=nivel,
        pct_aprobados_declarado=pct_aprobados_declarado,
        criterios=criterios,
        estudiantes=estudiantes,
        nombre_archivo_info=parsear_nombre_archivo(nombre_archivo),
    )
